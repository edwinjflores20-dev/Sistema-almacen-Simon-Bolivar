from flask import Flask, request, redirect, render_template_string
from openpyxl import Workbook, load_workbook
from datetime import date
from functools import reduce
import os


# ============================================================
# CONFIGURACIÓN DE LA APLICACIÓN
# ============================================================

app = Flask(__name__)

ARCHIVO_EXCEL = "almacen_hospital.xlsx"


# ============================================================
# CLASE PRODUCTO
# ENCAPSULAMIENTO CON GETTER Y SETTER
# ============================================================

class Producto:

    def __init__(self, id_producto, nombre, categoria, stock, minimo):
        self.id = id_producto
        self.nombre = nombre
        self.categoria = categoria
        self._stock = stock
        self.minimo = minimo

    # Getter
    @property
    def stock(self):
        return self._stock

    # Setter
    @stock.setter
    def stock(self, valor):

        if valor < 0:
            raise ValueError("El stock no puede ser negativo.")

        self._stock = valor


# ============================================================
# CREAR ARCHIVO EXCEL
# ============================================================

def crear_excel():

    if not os.path.exists(ARCHIVO_EXCEL):

        wb = Workbook()

        # Hoja Productos
        ws_productos = wb.active
        ws_productos.title = "Productos"

        ws_productos.append([
            "ID",
            "Nombre",
            "Categoria",
            "Stock",
            "Stock Minimo"
        ])

        ws_productos.append([
            101,
            "Paracetamol 500mg",
            "Medicamentos",
            100,
            30
        ])

        # Hoja Movimientos
        ws_movimientos = wb.create_sheet("Movimientos")

        ws_movimientos.append([
            "ID Movimiento",
            "ID Producto",
            "Tipo",
            "Cantidad",
            "Fecha",
            "Motivo"
        ])

        wb.save(ARCHIVO_EXCEL)
        wb.close()


# ============================================================
# OBTENER PRODUCTOS
# ============================================================

def obtener_productos():

    crear_excel()

    wb = load_workbook(ARCHIVO_EXCEL)
    ws = wb["Productos"]

    productos = []

    for fila in range(2, ws.max_row + 1):

        if ws.cell(fila, 1).value is not None:

            productos.append({
                "id": ws.cell(fila, 1).value,
                "nombre": ws.cell(fila, 2).value,
                "categoria": ws.cell(fila, 3).value,
                "stock": ws.cell(fila, 4).value,
                "minimo": ws.cell(fila, 5).value
            })

    wb.close()

    return productos


# ============================================================
# FUNCIÓN DE ORDEN SUPERIOR - REDUCE
# CALCULA EL STOCK TOTAL DEL SISTEMA
# ============================================================

def calcular_stock_total(productos):

    return reduce(
        lambda acumulado, producto:
        acumulado + producto["stock"],
        productos,
        0
    )


# ============================================================
# GUARDAR PRODUCTO
# ============================================================

def guardar_producto(
    id_producto,
    nombre,
    categoria,
    stock,
    minimo
):

    crear_excel()

    # Se utiliza la clase Producto
    producto = Producto(
        id_producto,
        nombre,
        categoria,
        stock,
        minimo
    )

    wb = load_workbook(ARCHIVO_EXCEL)

    ws = wb["Productos"]

    # Verificar que el ID no exista
    for fila in range(2, ws.max_row + 1):

        if ws.cell(fila, 1).value == id_producto:

            wb.close()

            return False

    # Guardar los datos
    ws.append([
        producto.id,
        producto.nombre,
        producto.categoria,
        producto.stock,
        producto.minimo
    ])

    wb.save(ARCHIVO_EXCEL)

    wb.close()

    return True


# ============================================================
# REGISTRAR MOVIMIENTO
# ============================================================

def registrar_movimiento(
    id_producto,
    tipo,
    cantidad,
    motivo
):

    crear_excel()

    wb = load_workbook(ARCHIVO_EXCEL)

    ws_productos = wb["Productos"]

    ws_movimientos = wb["Movimientos"]

    fila_producto = None

    # Buscar producto
    for fila in range(2, ws_productos.max_row + 1):

        if ws_productos.cell(fila, 1).value == id_producto:

            fila_producto = fila

            break

    if fila_producto is None:

        wb.close()

        return False, "Producto no encontrado."

    # Obtener stock actual
    stock_actual = ws_productos.cell(
        fila_producto,
        4
    ).value

    # Crear objeto Producto
    producto = Producto(
        ws_productos.cell(fila_producto, 1).value,
        ws_productos.cell(fila_producto, 2).value,
        ws_productos.cell(fila_producto, 3).value,
        stock_actual,
        ws_productos.cell(fila_producto, 5).value
    )

    # Movimiento de entrada
    if tipo == "Entrada":

        producto.stock = producto.stock + cantidad

    # Movimiento de salida
    elif tipo == "Salida":

        if cantidad > producto.stock:

            wb.close()

            return False, "Stock insuficiente."

        producto.stock = producto.stock - cantidad

    else:

        wb.close()

        return False, "Tipo de movimiento inválido."

    # Actualizar stock
    ws_productos.cell(
        fila_producto,
        4
    ).value = producto.stock

    # Obtener último ID
    ultimo_id = 0

    for fila in range(2, ws_movimientos.max_row + 1):

        valor = ws_movimientos.cell(
            fila,
            1
        ).value

        if valor is not None:

            if int(valor) > ultimo_id:

                ultimo_id = int(valor)

    nuevo_id = ultimo_id + 1

    # Registrar movimiento
    ws_movimientos.append([
        nuevo_id,
        id_producto,
        tipo,
        cantidad,
        date.today(),
        motivo
    ])

    wb.save(ARCHIVO_EXCEL)

    wb.close()

    return True, "Movimiento registrado correctamente."


# ============================================================
# OBTENER MOVIMIENTOS
# ============================================================

def obtener_movimientos():

    crear_excel()

    wb = load_workbook(ARCHIVO_EXCEL)

    ws = wb["Movimientos"]

    movimientos = []

    for fila in range(2, ws.max_row + 1):

        if ws.cell(fila, 1).value is not None:

            movimientos.append({
                "id": ws.cell(fila, 1).value,
                "producto": ws.cell(fila, 2).value,
                "tipo": ws.cell(fila, 3).value,
                "cantidad": ws.cell(fila, 4).value,
                "fecha": ws.cell(fila, 5).value,
                "motivo": ws.cell(fila, 6).value
            })

    wb.close()

    movimientos.reverse()

    return movimientos


# ============================================================
# INTERFAZ HTML
# ============================================================

HTML = """

<!DOCTYPE html>

<html lang="es">

<head>

<meta charset="UTF-8">

<meta name="viewport"
content="width=device-width, initial-scale=1.0">

<title>Sistema de Gestión de Almacén</title>

<style>

* {
    box-sizing: border-box;
    margin: 0;
    padding: 0;
}

body {
    font-family: Arial, Helvetica, sans-serif;
    background: #f4f7fa;
    color: #263238;
}

/* ==========================================================
   BARRA LATERAL
   ========================================================== */

.sidebar {

    position: fixed;

    left: 0;
    top: 0;

    width: 260px;
    height: 100vh;

    background: #123b5d;

    color: white;

    padding: 25px 15px;

    overflow-y: auto;
}

.logo {

    text-align: center;

    margin-bottom: 30px;
}

.logo h1 {

    font-size: 22px;

    margin-bottom: 8px;
}

.logo p {

    font-size: 13px;

    opacity: 0.8;
}

.menu a {

    display: block;

    text-decoration: none;

    color: white;

    padding: 14px 16px;

    margin-bottom: 8px;

    border-radius: 8px;

    transition: 0.3s;
}

.menu a:hover {

    background: #1d577f;
}

.menu a.active {

    background: #1d577f;
}

/* ==========================================================
   CONTENIDO PRINCIPAL
   ========================================================== */

.main {

    margin-left: 260px;

    padding: 30px;

    min-height: 100vh;
}

.header {

    margin-bottom: 25px;
}

.header h2 {

    color: #123b5d;

    margin-bottom: 6px;
}

.header p {

    color: #607d8b;
}

/* ==========================================================
   TARJETAS
   ========================================================== */

.cards {

    display: grid;

    grid-template-columns:
    repeat(4, 1fr);

    gap: 20px;

    margin-bottom: 25px;
}

.card {

    background: white;

    border-radius: 12px;

    padding: 22px;

    box-shadow:
    0 3px 12px rgba(0,0,0,0.08);
}

.card h3 {

    color: #607d8b;

    font-size: 14px;

    margin-bottom: 10px;
}

.card .numero {

    font-size: 30px;

    font-weight: bold;

    color: #123b5d;
}

.card.alerta .numero {

    color: #d32f2f;
}

/* ==========================================================
   TABLAS
   ========================================================== */

.table-container {

    background: white;

    padding: 20px;

    border-radius: 12px;

    overflow-x: auto;

    box-shadow:
    0 3px 12px rgba(0,0,0,0.08);
}

table {

    width: 100%;

    border-collapse: collapse;
}

th {

    background: #123b5d;

    color: white;

    padding: 13px;

    text-align: left;
}

td {

    padding: 12px;

    border-bottom: 1px solid #e0e0e0;
}

tr:hover {

    background: #f5f8fa;
}

/* ==========================================================
   ESTADOS
   ========================================================== */

.estado {

    display: inline-block;

    padding: 6px 10px;

    border-radius: 20px;

    font-size: 12px;

    font-weight: bold;
}

.estable {

    background: #e8f5e9;

    color: #2e7d32;
}

.critico {

    background: #ffebee;

    color: #c62828;
}

.entrada {

    color: #2e7d32;

    font-weight: bold;
}

.salida {

    color: #c62828;

    font-weight: bold;
}

/* ==========================================================
   FORMULARIOS
   ========================================================== */

.form-container {

    background: white;

    max-width: 800px;

    padding: 25px;

    border-radius: 12px;

    box-shadow:
    0 3px 12px rgba(0,0,0,0.08);
}

.form-group {

    margin-bottom: 18px;
}

.form-group label {

    display: block;

    font-weight: bold;

    margin-bottom: 7px;

    color: #37474f;
}

.form-group input,
.form-group select,
.form-group textarea {

    width: 100%;

    padding: 12px;

    border: 1px solid #cfd8dc;

    border-radius: 7px;

    font-size: 14px;
}

.form-group textarea {

    min-height: 100px;

    resize: vertical;
}

.btn {

    display: inline-block;

    border: none;

    background: #123b5d;

    color: white;

    padding: 12px 20px;

    border-radius: 7px;

    cursor: pointer;

    text-decoration: none;

    font-size: 14px;
}

.btn:hover {

    background: #1d577f;
}

.btn-secondary {

    background: #607d8b;
}

.btn-danger {

    background: #c62828;
}

/* ==========================================================
   ALERTAS
   ========================================================== */

.alert {

    padding: 15px;

    border-radius: 8px;

    margin-bottom: 20px;
}

.alert-error {

    background: #ffebee;

    color: #b71c1c;

    border-left: 5px solid #c62828;
}

.alert-success {

    background: #e8f5e9;

    color: #1b5e20;

    border-left: 5px solid #2e7d32;
}

/* ==========================================================
   TITULOS
   ========================================================== */

.section-title {

    color: #123b5d;

    margin-bottom: 18px;
}

.actions {

    display: flex;

    gap: 10px;

    flex-wrap: wrap;

    margin-bottom: 20px;
}

/* ==========================================================
   MENU MOVIL
   ========================================================== */

.mobile-menu {

    display: none;

    position: fixed;

    top: 15px;
    left: 15px;

    z-index: 1000;

    background: #123b5d;

    color: white;

    border: none;

    padding: 10px 14px;

    border-radius: 7px;

    cursor: pointer;
}

/* ==========================================================
   RESPONSIVE
   ========================================================== */

@media (max-width: 1000px) {

    .cards {

        grid-template-columns:
        repeat(2, 1fr);
    }
}

@media (max-width: 700px) {

    .sidebar {

        transform: translateX(-100%);

        transition: 0.3s;

        z-index: 999;
    }

    .sidebar.active {

        transform: translateX(0);
    }

    .main {

        margin-left: 0;

        padding: 70px 15px 20px;
    }

    .mobile-menu {

        display: block;
    }

    .cards {

        grid-template-columns: 1fr;
    }

    table {

        min-width: 650px;
    }

    .table-container {

        overflow-x: auto;
    }
}

</style>

</head>

<body>

<button
class="mobile-menu"
onclick="document.querySelector('.sidebar').classList.toggle('active')">

☰ Menú

</button>


<!-- ======================================================
     BARRA LATERAL
     ====================================================== -->

<aside class="sidebar">

    <div class="logo">

        <h1>🏥 Almacén</h1>

        <p>
            Sistema de Gestión
        </p>

    </div>

    <nav class="menu">

        <a href="/">
            🏠 Inicio
        </a>

        <a href="/productos">
            📦 Productos
        </a>

        <a href="/nuevo-producto">
            ➕ Nuevo producto
        </a>

        <a href="/movimiento">
            🔄 Registrar movimiento
        </a>

        <a href="/movimientos">
            📋 Historial
        </a>

        <a href="/alertas">
            ⚠️ Alertas
        </a>

    </nav>

</aside>


<!-- ======================================================
     CONTENIDO
     ====================================================== -->

<main class="main">

    {{ contenido | safe }}

</main>


</body>

</html>

"""


# ============================================================
# FUNCIÓN PARA MOSTRAR PÁGINAS
# ============================================================

def pagina(contenido):

    return render_template_string(
        HTML,
        contenido=contenido
    )


# ============================================================
# INICIO / DASHBOARD
# ============================================================

@app.route("/")
def inicio():

    productos = obtener_productos()

    total_productos = len(productos)

    # REDUCE aplicado a la colección de productos
    total_stock = calcular_stock_total(productos)

    stock_critico = 0

    for producto in productos:

        if producto["stock"] <= producto["minimo"]:

            stock_critico += 1

    movimientos = obtener_movimientos()

    total_movimientos = len(movimientos)

    contenido = f"""

    <div class="header">

        <h2>
            Sistema de Gestión de Almacén
        </h2>

        <p>
            Hospital Simón Bolívar
        </p>

    </div>


    <div class="cards">

        <div class="card">

            <h3>
                Productos registrados
            </h3>

            <div class="numero">
                {total_productos}
            </div>

        </div>


        <div class="card">

            <h3>
                Stock total
            </h3>

            <div class="numero">
                {total_stock}
            </div>

        </div>


        <div class="card alerta">

            <h3>
                Stock crítico
            </h3>

            <div class="numero">
                {stock_critico}
            </div>

        </div>


        <div class="card">

            <h3>
                Movimientos
            </h3>

            <div class="numero">
                {total_movimientos}
            </div>

        </div>

    </div>


    <div class="table-container">

        <h3 class="section-title">
            Estado del inventario
        </h3>

        <table>

            <thead>

                <tr>

                    <th>ID</th>

                    <th>Producto</th>

                    <th>Categoría</th>

                    <th>Stock</th>

                    <th>Mínimo</th>

                    <th>Estado</th>

                </tr>

            </thead>

            <tbody>
    """

    if not productos:

        contenido += """

            <tr>

                <td colspan="6">
                    No hay productos registrados.
                </td>

            </tr>

        """

    else:

        for producto in productos:

            if producto["stock"] <= producto["minimo"]:

                estado = """

                <span class="estado critico">
                    Crítico
                </span>

                """

            else:

                estado = """

                <span class="estado estable">
                    Estable
                </span>

                """

            contenido += f"""

                <tr>

                    <td>
                        {producto["id"]}
                    </td>

                    <td>
                        {producto["nombre"]}
                    </td>

                    <td>
                        {producto["categoria"]}
                    </td>

                    <td>
                        {producto["stock"]}
                    </td>

                    <td>
                        {producto["minimo"]}
                    </td>

                    <td>
                        {estado}
                    </td>

                </tr>

            """

    contenido += """

            </tbody>

        </table>

    </div>

    """

    return pagina(contenido)


# ============================================================
# LISTA DE PRODUCTOS
# ============================================================

@app.route("/productos")
def productos():

    lista = obtener_productos()

    contenido = """

    <div class="header">

        <h2>
            Productos
        </h2>

        <p>
            Productos registrados en el almacén.
        </p>

    </div>


    <div class="actions">

        <a
        href="/nuevo-producto"
        class="btn">

            ➕ Nuevo producto

        </a>

    </div>


    <div class="table-container">

        <table>

            <thead>

                <tr>

                    <th>ID</th>

                    <th>Nombre</th>

                    <th>Categoría</th>

                    <th>Stock</th>

                    <th>Mínimo</th>

                    <th>Estado</th>

                </tr>

            </thead>

            <tbody>

    """

    if not lista:

        contenido += """

            <tr>

                <td colspan="6">

                    No hay productos registrados.

                </td>

            </tr>

        """

    else:

        for producto in lista:

            if producto["stock"] <= producto["minimo"]:

                estado = """

                <span class="estado critico">
                    Crítico
                </span>

                """

            else:

                estado = """

                <span class="estado estable">
                    Estable
                </span>

                """

            contenido += f"""

                <tr>

                    <td>
                        {producto["id"]}
                    </td>

                    <td>
                        {producto["nombre"]}
                    </td>

                    <td>
                        {producto["categoria"]}
                    </td>

                    <td>
                        {producto["stock"]}
                    </td>

                    <td>
                        {producto["minimo"]}
                    </td>

                    <td>
                        {estado}
                    </td>

                </tr>

            """

    contenido += """

            </tbody>

        </table>

    </div>

    """

    return pagina(contenido)


# ============================================================
# FORMULARIO NUEVO PRODUCTO
# ============================================================

@app.route("/nuevo-producto")
def nuevo_producto():

    contenido = """

    <div class="header">

        <h2>
            Nuevo producto
        </h2>

        <p>
            Registrar un nuevo producto en el almacén.
        </p>

    </div>


    <div class="form-container">

        <form
        action="/guardar-producto"
        method="POST">


            <div class="form-group">

                <label>
                    ID del producto
                </label>

                <input
                type="number"
                name="id"
                required
                min="1">

            </div>


            <div class="form-group">

                <label>
                    Nombre del producto
                </label>

                <input
                type="text"
                name="nombre"
                required>

            </div>


            <div class="form-group">

                <label>
                    Categoría
                </label>

                <input
                type="text"
                name="categoria"
                required>

            </div>


            <div class="form-group">

                <label>
                    Stock inicial
                </label>

                <input
                type="number"
                name="stock"
                required
                min="0">

            </div>


            <div class="form-group">

                <label>
                    Stock mínimo
                </label>

                <input
                type="number"
                name="minimo"
                required
                min="0">

            </div>


            <button
            type="submit"
            class="btn">

                Guardar producto

            </button>


            <a
            href="/productos"
            class="btn btn-secondary">

                Cancelar

            </a>

        </form>

    </div>

    """

    return pagina(contenido)


# ============================================================
# GUARDAR NUEVO PRODUCTO
# MANEJO DE ERRORES Y EXCEPCIONES
# ============================================================

@app.route(
    "/guardar-producto",
    methods=["POST"]
)
def guardar_nuevo_producto():

    try:

        id_producto = int(
            request.form["id"]
        )

        nombre = request.form[
            "nombre"
        ].strip()

        categoria = request.form[
            "categoria"
        ].strip()

        stock = int(
            request.form["stock"]
        )

        minimo = int(
            request.form["minimo"]
        )

    except (ValueError, KeyError):

        contenido = """

        <div class="alert alert-error">

            <strong>Error:</strong>

            Los datos ingresados
            no son válidos.

        </div>

        <a
        href="/nuevo-producto"
        class="btn">

            Volver

        </a>

        """

        return pagina(contenido)


    # ========================================================
    # VALIDACIÓN DE DATOS
    # ========================================================

    if stock < 0 or minimo < 0:

        contenido = """

        <div class="alert alert-error">

            <strong>Error:</strong>

            El stock y el stock mínimo
            no pueden ser negativos.

        </div>

        <a
        href="/nuevo-producto"
        class="btn">

            Volver

        </a>

        """

        return pagina(contenido)


    if not nombre or not categoria:

        contenido = """

        <div class="alert alert-error">

            <strong>Error:</strong>

            El nombre y la categoría
            son obligatorios.

        </div>

        <a
        href="/nuevo-producto"
        class="btn">

            Volver

        </a>

        """

        return pagina(contenido)


    resultado = guardar_producto(
        id_producto,
        nombre,
        categoria,
        stock,
        minimo
    )


    if resultado:

        return redirect("/productos")


    contenido = """

    <div class="alert alert-error">

        <strong>Error:</strong>

        Ya existe un producto
        con ese ID.

    </div>


    <a
    href="/nuevo-producto"
    class="btn">

        Volver

    </a>

    """

    return pagina(contenido)


# ============================================================
# FORMULARIO DE MOVIMIENTOS
# ============================================================

@app.route("/movimiento")
def movimiento():

    productos = obtener_productos()

    contenido = """

    <div class="header">

        <h2>
            Registrar movimiento
        </h2>

        <p>
            Registre entradas o salidas
            de productos.
        </p>

    </div>


    <div class="form-container">

        <form
        action="/guardar-movimiento"
        method="POST">


            <div class="form-group">

                <label>
                    Producto
                </label>

                <select
                name="id_producto"
                required>

                    <option value="">
                        Seleccione un producto
                    </option>

    """

    for producto in productos:

        contenido += f"""

                    <option value="{producto["id"]}">

                        {producto["id"]}
                        -
                        {producto["nombre"]}
                        |
                        Stock:
                        {producto["stock"]}

                    </option>

        """

    contenido += """

                </select>

            </div>


            <div class="form-group">

                <label>
                    Tipo de movimiento
                </label>

                <select
                name="tipo"
                required>

                    <option value="">
                        Seleccione
                    </option>

                    <option value="Entrada">
                        Entrada
                    </option>

                    <option value="Salida">
                        Salida
                    </option>

                </select>

            </div>


            <div class="form-group">

                <label>
                    Cantidad
                </label>

                <input
                type="number"
                name="cantidad"
                required
                min="1">

            </div>


            <div class="form-group">

                <label>
                    Motivo
                </label>

                <textarea
                name="motivo"
                required
                placeholder="Ingrese el motivo del movimiento"></textarea>

            </div>


            <button
            type="submit"
            class="btn">

                Registrar movimiento

            </button>


            <a
            href="/"
            class="btn btn-secondary">

                Cancelar

            </a>

        </form>

    </div>

    """

    return pagina(contenido)


# ============================================================
# GUARDAR MOVIMIENTO
# MANEJO DE ERRORES
# ============================================================

@app.route(
    "/guardar-movimiento",
    methods=["POST"]
)
def guardar_nuevo_movimiento():

    try:

        id_producto = int(
            request.form["id_producto"]
        )

        tipo = request.form["tipo"]

        cantidad = int(
            request.form["cantidad"]
        )

        motivo = request.form[
            "motivo"
        ].strip()

    except (ValueError, KeyError):

        contenido = """

        <div class="alert alert-error">

            <strong>Error:</strong>

            Los datos ingresados
            no son válidos.

        </div>


        <a
        href="/movimiento"
        class="btn">

            Volver

        </a>

        """

        return pagina(contenido)


    if cantidad <= 0:

        contenido = """

        <div class="alert alert-error">

            <strong>Error:</strong>

            La cantidad debe ser
            mayor que cero.

        </div>


        <a
        href="/movimiento"
        class="btn">

            Volver

        </a>

        """

        return pagina(contenido)


    if tipo not in ["Entrada", "Salida"]:

        contenido = """

        <div class="alert alert-error">

            <strong>Error:</strong>

            Tipo de movimiento inválido.

        </div>


        <a
        href="/movimiento"
        class="btn">

            Volver

        </a>

        """

        return pagina(contenido)


    if not motivo:

        contenido = """

        <div class="alert alert-error">

            <strong>Error:</strong>

            El motivo es obligatorio.

        </div>


        <a
        href="/movimiento"
        class="btn">

            Volver

        </a>

        """

        return pagina(contenido)


    resultado, mensaje = registrar_movimiento(
        id_producto,
        tipo,
        cantidad,
        motivo
    )


    if resultado:

        return redirect("/productos")


    contenido = f"""

    <div class="alert alert-error">

        <strong>Error:</strong>

        {mensaje}

    </div>


    <a
    href="/movimiento"
    class="btn">

        Volver

    </a>

    """

    return pagina(contenido)


# ============================================================
# HISTORIAL DE MOVIMIENTOS
# ============================================================

@app.route("/movimientos")
def movimientos():

    lista = obtener_movimientos()

    contenido = """

    <div class="header">

        <h2>
            Historial de movimientos
        </h2>

        <p>
            Registro de entradas y salidas
            del almacén.
        </p>

    </div>


    <div class="table-container">

        <table>

            <thead>

                <tr>

                    <th>ID Movimiento</th>

                    <th>ID Producto</th>

                    <th>Tipo</th>

                    <th>Cantidad</th>

                    <th>Fecha</th>

                    <th>Motivo</th>

                </tr>

            </thead>

            <tbody>

    """


    if not lista:

        contenido += """

            <tr>

                <td colspan="6">

                    No existen movimientos registrados.

                </td>

            </tr>

        """

    else:

        for movimiento in lista:

            clase = (
                "entrada"
                if movimiento["tipo"] == "Entrada"
                else "salida"
            )

            fecha = movimiento["fecha"]

            contenido += f"""

                <tr>

                    <td>
                        {movimiento["id"]}
                    </td>

                    <td>
                        {movimiento["producto"]}
                    </td>

                    <td class="{clase}">
                        {movimiento["tipo"]}
                    </td>

                    <td>
                        {movimiento["cantidad"]}
                    </td>

                    <td>
                        {fecha}
                    </td>

                    <td>
                        {movimiento["motivo"]}
                    </td>

                </tr>

            """


    contenido += """

            </tbody>

        </table>

    </div>

    """

    return pagina(contenido)


# ============================================================
# ALERTAS DE STOCK
# ============================================================

@app.route("/alertas")
def alertas():

    productos = obtener_productos()

    productos_criticos = []

    for producto in productos:

        if producto["stock"] <= producto["minimo"]:

            productos_criticos.append(producto)


    contenido = """

    <div class="header">

        <h2>
            Alertas de stock
        </h2>

        <p>
            Productos que necesitan reposición.
        </p>

    </div>

    """


    if not productos_criticos:

        contenido += """

        <div class="alert alert-success">

            <strong>Todo correcto.</strong>

            No existen productos
            con stock crítico.

        </div>

        """

    else:

        contenido += """

        <div class="alert alert-error">

            Existen productos cuyo stock
            está por debajo o igual
            al mínimo establecido.

        </div>


        <div class="table-container">

            <table>

                <thead>

                    <tr>

                        <th>ID</th>

                        <th>Producto</th>

                        <th>Categoría</th>

                        <th>Stock</th>

                        <th>Mínimo</th>

                        <th>Estado</th>

                    </tr>

                </thead>

                <tbody>

        """


        for producto in productos_criticos:

            contenido += f"""

                    <tr>

                        <td>
                            {producto["id"]}
                        </td>

                        <td>
                            {producto["nombre"]}
                        </td>

                        <td>
                            {producto["categoria"]}
                        </td>

                        <td>
                            {producto["stock"]}
                        </td>

                        <td>
                            {producto["minimo"]}
                        </td>

                        <td>

                            <span class="estado critico">

                                Crítico

                            </span>

                        </td>

                    </tr>

            """


        contenido += """

                </tbody>

            </table>

        </div>

        """


    return pagina(contenido)


# ============================================================
# EJECUTAR APLICACIÓN
# ============================================================

if __name__ == "__main__":

    crear_excel()

    print("=" * 50)

    print(
        "SISTEMA DE GESTIÓN DE ALMACÉN HOSPITALARIO"
    )

    print("=" * 50)

    print(
        "Servidor local:"
    )

    print(
        "http://127.0.0.1:5000"
    )

    print(
        "Archivo Excel:",
        ARCHIVO_EXCEL
    )

    print("=" * 50)

    app.run(
        host="0.0.0.0",
        port=5000,
        debug=True
    )